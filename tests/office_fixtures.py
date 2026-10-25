"""Tiny valid OOXML files for tests. Built with python-docx / python-pptx."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptInches

from noteagent.acquire.media.png import encode_png

LONG_BODY = "KL散度是两个概率分布之间的差异度量。" * 20
WORD_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PPT_MIME = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def _png(width: int = 80, height: int = 80) -> bytes:
    pixels = bytes((200, 16, 16)) * (width * height)
    return encode_png(width, height, pixels, color_type=2)


def build_docx(
    *,
    heading: str = "相对熵",
    paragraphs: list[str] | None = None,
    table: list[list[str]] | None = None,
    with_image: bool = False,
    title: str = "",
) -> bytes:
    document = Document()
    if title:
        document.core_properties.title = title
    if heading:
        document.add_heading(heading, level=1)
    for text in paragraphs if paragraphs is not None else [LONG_BODY]:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for col_index, cell in enumerate(row):
                grid.cell(row_index, col_index).text = cell
    if with_image:
        document.add_paragraph().add_run().add_picture(
            BytesIO(_png()),
            width=Inches(1.2),
        )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pptx(
    *,
    slides: list[str] | None = None,
    with_image: bool = False,
    title: str = "",
) -> bytes:
    presentation = Presentation()
    if title:
        presentation.core_properties.title = title
    blank = (
        presentation.slide_layouts[6]
        if len(presentation.slide_layouts) > 6
        else presentation.slide_layouts[0]
    )
    texts = slides if slides is not None else [LONG_BODY, "第二页补充说明。" * 8]
    for index, text in enumerate(texts):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(
            PptInches(0.5),
            PptInches(0.5),
            PptInches(9),
            PptInches(2.5),
        )
        box.text_frame.text = text
        if with_image and index == 0:
            slide.shapes.add_picture(
                BytesIO(_png()),
                PptInches(1),
                PptInches(3.2),
                width=PptInches(1.5),
            )
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


HELLO_DOCX = build_docx()
HELLO_PPTX = build_pptx()
IMAGE_DOCX = build_docx(with_image=True)
TABLE_DOCX = build_docx(
    paragraphs=["正文段落。"],
    table=[["符号", "含义"], ["D_KL", "相对熵"]],
)
GARBAGE_OFFICE = b"PK\x03\x04not-a-real-office"
