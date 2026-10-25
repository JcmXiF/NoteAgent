"""Office originals: immutable OOXML bytes + derived HTML leaves."""

from __future__ import annotations

import html as html_lib
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from noteagent.acquire.adapters.base import SourceKind
from noteagent.acquire.derive import DerivedBlock
from noteagent.acquire.media.captions import CaptionStore
from noteagent.acquire.media.extract import (
    ExtractedImage,
    apply_source_cap,
    captions_for_images,
    image_from_bytes,
    image_marker,
    join_page_text,
    media_extract_enabled,
    persist_extracted_images,
)
from noteagent.acquire.render import wrap_article_document
from noteagent.project.layout import ProjectLayout

OFFICE_CAPTURE_METHOD = "office-ooxml"
MAX_OFFICE_BYTES = 20_000_000
OFFICE_ACCEPT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document,"
    "application/vnd.openxmlformats-officedocument.presentationml.presentation,"
    "application/octet-stream;q=0.9,*/*;q=0.8"
)
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


class OfficeDeriveError(RuntimeError):
    """Could not read this OOXML package. Raw bytes are unchanged."""


@dataclass(frozen=True, slots=True)
class OfficeDerived:
    blocks: list[DerivedBlock]
    html: str
    subtype: str


class OfficeSourceAdapter:
    """Persist docx/pptx under raw/office/; derive stamped HTML leaves."""

    kind = SourceKind.OFFICE

    def derive(
        self,
        path: Path,
        *,
        layout: ProjectLayout | None = None,
        identity: str = "",
    ) -> list[DerivedBlock]:
        return self.derive_bytes(
            Path(path).read_bytes(),
            filename=Path(path).name,
            layout=layout,
            identity=identity,
        ).blocks

    def derive_bytes(
        self,
        data: bytes,
        *,
        filename: str = "",
        layout: ProjectLayout | None = None,
        identity: str = "",
        title: str = "",
        source_url: str = "",
        capture_method: str = OFFICE_CAPTURE_METHOD,
        fidelity: str = "ok",
    ) -> OfficeDerived:
        subtype = sniff_office_package(data)
        if subtype is None:
            raise OfficeDeriveError("not a docx/pptx package")
        store = CaptionStore.for_layout(layout) if layout is not None else CaptionStore.memory()
        try:
            if subtype == "pptx":
                rows = _pptx_rows(data)
            else:
                rows = _docx_rows(data)
        except (OfficeDeriveError, ValueError, KeyError, OSError) as err:
            raise OfficeDeriveError(str(err)) from err
        except Exception as err:
            raise OfficeDeriveError(str(err)) from err

        flat = [image for row in rows for image in row.images]
        capped = apply_source_cap(flat) if media_extract_enabled() else []
        by_id: dict[int, list[ExtractedImage]] = {}
        for image in capped:
            if image.skip_reason or image.page is None:
                continue
            by_id.setdefault(image.page, []).append(image)
        if layout is not None and identity:
            persist_extracted_images(layout, identity, capped, store=store)

        blocks: list[DerivedBlock] = []
        html_parts: list[str] = []
        for row in rows:
            captions = captions_for_images(by_id.get(row.index, []), store)
            markers = [image_marker(text) for text in captions]
            text = join_page_text(row.text, markers, empty="")
            if not text:
                continue
            blocks.append(DerivedBlock(id=row.block_id, tag=row.tag, text=text))
            html_parts.append(_leaf_html(row.block_id, row.tag, text, row.table_rows))
        if layout is not None:
            store.save()
        body = "\n".join(html_parts)
        html = wrap_article_document(
            title=title or Path(filename).stem or "office",
            source_url=source_url,
            capture_method=capture_method,
            fidelity=fidelity,
            body_html=body,
        )
        return OfficeDerived(blocks=blocks, html=html, subtype=subtype)

    def preview(self, path: Path) -> dict[str, str]:
        return {
            "kind": self.kind.value,
            "renderer": "html-sandbox",
            "path": Path(path).as_posix(),
        }


@dataclass(frozen=True, slots=True)
class _Row:
    index: int
    block_id: str
    tag: str
    text: str
    images: list[ExtractedImage]
    table_rows: tuple[tuple[str, ...], ...] = ()


def sniff_office_package(data: bytes) -> str | None:
    """Return ``docx`` / ``pptx`` when the ZIP is a Word or PowerPoint package."""
    if not data.startswith(b"PK"):
        return None
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = set(archive.namelist())
    except (BadZipFile, OSError):
        return None
    if "word/document.xml" in names:
        return "docx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return None


def looks_like_office(data: bytes) -> bool:
    return sniff_office_package(data) is not None


def _docx_rows(data: bytes) -> list[_Row]:
    document = Document(BytesIO(data))
    rows: list[_Row] = []
    index = 0
    for item in _iter_docx_blocks(document):
        index += 1
        if isinstance(item, Table):
            table_rows = _table_cells(item)
            text = "\n".join(" | ".join(cells) for cells in table_rows if any(cells))
            images = _element_images(item._tbl, document.part, index)
            if not text and not images:
                continue
            rows.append(
                _Row(
                    index=index,
                    block_id=f"s{len(rows) + 1}",
                    tag="table",
                    text=text,
                    images=images,
                    table_rows=tuple(table_rows),
                )
            )
            continue
        tag = _heading_tag(item)
        text = " ".join(item.text.split())
        images = _paragraph_images(item, index)
        if not text and not images:
            continue
        rows.append(
            _Row(
                index=index,
                block_id=f"s{len(rows) + 1}",
                tag=tag,
                text=text,
                images=images,
            )
        )
    return rows


def _pptx_rows(data: bytes) -> list[_Row]:
    presentation = Presentation(BytesIO(data))
    rows: list[_Row] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        images: list[ExtractedImage] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    compact = " ".join(paragraph.text.split())
                    if compact:
                        parts.append(compact)
            images.extend(_shape_images(shape, number))
        text = " ".join(parts)
        if not text and not images:
            continue
        rows.append(
            _Row(
                index=number,
                block_id=f"slide-{number}",
                tag="section",
                text=text,
                images=images,
            )
        )
    return rows


def _iter_docx_blocks(document: DocumentObject):
    body = document.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _heading_tag(paragraph: Paragraph) -> str:
    style = paragraph.style
    name = (style.name or "") if style is not None else ""
    lowered = name.strip().lower()
    if lowered.startswith("heading"):
        suffix = lowered.replace("heading", "").strip()
        if suffix.isdigit() and 1 <= int(suffix) <= 4:
            return f"h{int(suffix)}"
    if lowered.startswith("标题"):
        suffix = lowered.replace("标题", "").strip()
        if suffix.isdigit() and 1 <= int(suffix) <= 4:
            return f"h{int(suffix)}"
    return "p"


def _table_cells(table: Table) -> tuple[tuple[str, ...], ...]:
    rows: list[tuple[str, ...]] = []
    for row in table.rows:
        cells = tuple(" ".join(cell.text.split()) for cell in row.cells)
        rows.append(cells)
    return tuple(rows)


def _paragraph_images(paragraph: Paragraph, index: int) -> list[ExtractedImage]:
    return _element_images(paragraph._element, paragraph.part, index)


def _element_images(element: object, part: object, index: int) -> list[ExtractedImage]:
    found: list[ExtractedImage] = []
    if not media_extract_enabled():
        return found
    related = getattr(part, "related_parts", {})
    for blip in element.findall(f".//{_A_BLIP}"):
        rid = blip.get(_R_EMBED)
        if not rid:
            continue
        try:
            blob = related[rid].blob
        except (KeyError, AttributeError):
            continue
        image = image_from_bytes(blob, page=index)
        if image is not None:
            found.append(image)
    return found


def _shape_images(shape: object, index: int) -> list[ExtractedImage]:
    if not media_extract_enabled():
        return []
    try:
        shape_type = shape.shape_type
    except (AttributeError, ValueError):
        return []
    if shape_type != MSO_SHAPE_TYPE.PICTURE:
        return []
    try:
        blob = shape.image.blob
    except (AttributeError, ValueError, OSError):
        return []
    image = image_from_bytes(blob, page=index)
    return [image] if image is not None else []


def _leaf_html(
    block_id: str,
    tag: str,
    text: str,
    table_rows: tuple[tuple[str, ...], ...],
) -> str:
    safe_id = html_lib.escape(block_id, quote=True)
    if tag == "table":
        body_rows = table_rows or (tuple(text.split(" | ")),)
        rows_html: list[str] = []
        for cells in body_rows:
            tds = "".join(f"<td>{html_lib.escape(cell)}</td>" for cell in cells)
            rows_html.append(f"<tr>{tds}</tr>")
        extra = ""
        if "[图" in text and not any("[图" in cell for row in body_rows for cell in row):
            extra = f"<tr><td>{html_lib.escape(text)}</td></tr>"
        return (
            f'<table id="na-{safe_id}" data-na-block="{safe_id}">'
            f"{''.join(rows_html)}{extra}</table>"
        )
    if tag not in {"p", "h1", "h2", "h3", "h4", "section"}:
        tag = "p"
    return (
        f'<{tag} id="na-{safe_id}" data-na-block="{safe_id}">'
        f"{html_lib.escape(text)}</{tag}>"
    )
